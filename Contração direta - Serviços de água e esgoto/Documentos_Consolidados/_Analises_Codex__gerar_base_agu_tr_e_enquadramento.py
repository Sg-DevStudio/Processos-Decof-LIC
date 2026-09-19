from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import RGBColor, Pt


BASE = Path(r"C:\Users\Samuel Gomes\Desktop\Processos - Decof_LIC\Contração direta - Serviços de água e esgoto")
TEMPLATE = BASE / "Modelos_Oficiais_AGU" / "01_Modelo_AGU_TR_Servicos_e_Obras_Maio_2026.docx"
OUT = BASE / "Documentos_da_Instrução"
OUT.mkdir(exist_ok=True)

BLUE = "0000FF"
RED = "FF0000"
GREEN = "008000"
GRAY = "808080"


def wipe_body(document):
    body = document._element.body
    section_props = body.sectPr
    for child in list(body):
        if child is not section_props:
            body.remove(child)


def set_run(run, text, color=None, bold=None, italic=None, strike=False, size=None):
    run.text = text
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if strike:
        run.font.strike = True
    if size:
        run.font.size = Pt(size)
    return run


def add_para(document, parts=(), style="Normal", align=None, before=0, after=90):
    p = document.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    for part in parts:
        if isinstance(part, str):
            set_run(p.add_run(), part)
        else:
            text, color, bold, italic, strike = part
            set_run(p.add_run(), text, color, bold, italic, strike)
    return p


def heading(document, title):
    return add_para(document, [(title, None, True, False, False)], style="Nivel 01", before=170, after=80)


def optional(document, text):
    return add_para(document, [(text, BLUE, False, False, False)], style="Nível 02", after=80)


def rationale(document, text):
    return add_para(document, [("JUSTIFICATIVA DE ALTERAÇÃO: " + text, GRAY, False, True, False)], style="Normal", after=110)


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def fill_table(document, table_template):
    # Keeps the original AGU table grid, widths and cell formatting.
    document._element.body.insert(len(document._element.body) - 1, deepcopy(table_template))
    table = document.tables[-1]
    while len(table.rows) > 3:
        table._tbl.remove(table.rows[-1]._tr)

    headers = ["ITEM", "ESPECIFICAÇÃO", "CATSER", "UNIDADE", "QUANTIDADE", "VALOR UNITÁRIO", "VALOR TOTAL"]
    for cell, value in zip(table.rows[0].cells, headers):
        p = clear_cell(cell)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(), value, bold=True, size=8)

    values = [
        "1",
        "Fornecimento contínuo de água tratada e coleta/tratamento de esgoto sanitário para as unidades consumidoras do Complexo de São Cristóvão.",
        "[preencher]",
        "mês",
        "12",
        "[preencher após ETP]",
        "R$ 1.180.116,40",
    ]
    for cell, value in zip(table.rows[1].cells, values):
        p = clear_cell(cell)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if value in {"1", "mês", "12"} else WD_ALIGN_PARAGRAPH.LEFT
        color = BLUE if value.startswith("[") else RED
        set_run(p.add_run(), value, color=color, size=8)

    merged = table.rows[2].cells[0].merge(table.rows[2].cells[5])
    p = clear_cell(merged)
    set_run(p.add_run(), "Estimativa anual baseada na Nota Técnica de Justificativa de Preços e no histórico de faturas.", color=GRAY, italic=True, size=8)
    p = clear_cell(table.rows[2].cells[6])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(), "R$ 1.180.116,40", color=RED, bold=True, size=8)
    return table


def add_cover(document, subtitle):
    add_para(document, [("COLÉGIO PEDRO II", BLUE, True, False, False)], align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=35)
    add_para(document, [("(Processo Administrativo nº ", None, False, False, False), ("[preencher NUP]", BLUE, False, False, False), (")", None, False, False, False)], align=WD_ALIGN_PARAGRAPH.CENTER, after=80)
    add_para(document, [("TERMO DE REFERÊNCIA", None, True, False, False)], align=WD_ALIGN_PARAGRAPH.CENTER, before=60, after=35)
    add_para(document, [(subtitle, BLUE, True, False, False)], align=WD_ALIGN_PARAGRAPH.CENTER, after=120)


def build_tr():
    output = OUT / "DOC ÁGUA - 07 — Termo de Referência padrão AGU Maio 2026 (base ETP).docx"
    shutil.copy2(TEMPLATE, output)
    document = Document(output)
    table_template = deepcopy(document.tables[0]._tbl)
    wipe_body(document)

    add_cover(document, "MINUTA-BASE — adequar ao ETP antes da assinatura")
    rationale(document, "Minuta reconstruída sobre o modelo oficial da AGU de maio de 2026. Texto em vermelho é inclusão; azul indica lacuna a preencher; verde representa adaptação do modelo; cinza apresenta justificativa de alteração.")

    heading(document, "CONDIÇÕES GERAIS DA CONTRATAÇÃO")
    add_para(document, [
        ("Contratação de ", GREEN, False, False, False),
        ("serviços continuados", GREEN, False, False, False),
        (" de fornecimento de água tratada e coleta/tratamento de esgoto sanitário, a serem prestados pela concessionária responsável pela área das unidades consumidoras do Complexo de São Cristóvão do Colégio Pedro II, nos termos da tabela abaixo e das condições deste instrumento.", RED, False, False, False),
    ], style="Nível 02")
    fill_table(document, table_template)
    add_para(document, [("O objeto é caracterizado como serviço comum e continuado, conforme justificativa a ser consolidada no Estudo Técnico Preliminar.", GREEN, False, False, False)], style="Nível 02")
    optional(document, "O prazo de vigência será definido conforme o ETP, a análise jurídica e a natureza de serviço público monopolizado, com a demonstração anual de créditos orçamentários quando aplicável.")
    rationale(document, "A minuta não fixa prazo indeterminado antes de o ETP e a análise jurídica confirmarem a hipótese do art. 109 da Lei nº 14.133/2021.")

    heading(document, "FUNDAMENTAÇÃO E DESCRIÇÃO DA NECESSIDADE DA CONTRATAÇÃO")
    add_para(document, [("A contratação visa assegurar o abastecimento contínuo de água potável e a coleta/tratamento de esgoto sanitário necessários ao funcionamento das atividades pedagógicas, administrativas, sanitárias e de apoio do Complexo de São Cristóvão.", RED, False, False, False)], style="Nível 02")
    optional(document, "A fundamentação da contratação, dos quantitativos, das matrículas consumidoras, das economias e da previsão no Plano de Contratações Anual será pormenorizada no ETP e nos documentos anexos.")

    heading(document, "DESCRIÇÃO DA SOLUÇÃO COMO UM TODO CONSIDERADO O CICLO DE VIDA DO OBJETO")
    add_para(document, [("A solução consiste na prestação, pela concessionária local, dos serviços de fornecimento de água tratada e de coleta/tratamento de esgoto nas unidades consumidoras identificadas no Anexo I, com faturamento mensal conforme o consumo efetivamente medido e as tarifas reguladas vigentes.", RED, False, False, False)], style="Nível 02")

    heading(document, "REQUISITOS DA CONTRATAÇÃO")
    add_para(document, [("A contratada deverá prestar os serviços em conformidade com a regulação aplicável, manter a continuidade e a qualidade do abastecimento, emitir faturas individualizadas por matrícula e disponibilizar os canais de atendimento para ocorrências, interrupções e esclarecimentos.", RED, False, False, False)], style="Nível 02")
    add_para(document, [("Não será admitida a subcontratação do objeto contratual, pois a execução decorre da concessão e da atuação exclusiva da concessionária na área atendida.", GREEN, False, False, False)], style="Nível 02")
    add_para(document, [("Não haverá exigência de garantia contratual, pelas características do objeto e pela forma de cobrança regulada, sem prejuízo da validação no ETP.", RED, False, False, False)], style="Nível 02")

    heading(document, "MODELO DE EXECUÇÃO DO OBJETO")
    add_para(document, [("Os serviços serão prestados continuamente nas unidades e matrículas indicadas no Anexo I, mediante disponibilidade da rede, medição do consumo e faturamento mensal. A execução compreende o fornecimento de água tratada e a coleta/tratamento de esgoto sanitário, nos limites da infraestrutura disponível e das normas regulatórias.", RED, False, False, False)], style="Nível 02")
    optional(document, "Preencher no Anexo I os endereços, as matrículas, as economias e demais dados de cada unidade consumidora.")

    heading(document, "MODELO DE GESTÃO DO CONTRATO")
    add_para(document, [("A execução será acompanhada e fiscalizada por servidores formalmente designados pelo Colégio Pedro II, nos termos do art. 117 da Lei nº 14.133/2021.", RED, False, False, False)], style="Nível 02")
    optional(document, "Gestor: [nome, SIAPE e unidade]. Fiscal técnico/requisitante: [nome, SIAPE e unidade]. Fiscal substituto: [nome, SIAPE e unidade].")

    heading(document, "CRITÉRIOS DE MEDIÇÃO E DE PAGAMENTO")
    add_para(document, [("A medição será mensal, com base no consumo efetivamente registrado nas matrículas constantes do Anexo I e nas tarifas públicas/regulatórias vigentes. O pagamento será efetuado após o ateste da fatura pelo fiscal competente, observado o valor efetivamente devido no período.", RED, False, False, False)], style="Nível 02")
    add_para(document, [("A estimativa anual não estabelece consumo mínimo ou pagamento fixo; constitui teto estimativo para a contratação. Reajustes tarifários homologados deverão ser comprovados e conferidos nos autos.", RED, False, False, False)], style="Nível 02")

    heading(document, "FORMA E CRITÉRIOS DE SELEÇÃO DO FORNECEDOR")
    add_para(document, [("A seleção ocorrerá por contratação direta, mediante inexigibilidade de licitação, com fundamento no art. 74, inciso I, da Lei nº 14.133/2021, condicionada à comprovação documental da inviabilidade de competição e da exclusividade territorial da concessionária para as unidades do CPII.", RED, False, False, False)], style="Nível 02")
    add_para(document, [("Antes da contratação, deverão ser juntados o Formulário de Enquadramento, o documento de concessão e/ou declaração atual que vincule as matrículas/endereço à concessionária, a habilitação e a regularidade aplicáveis, a razão da escolha e a justificativa do preço.", RED, False, False, False)], style="Nível 02")

    heading(document, "ESTIMATIVAS DO VALOR DA CONTRATAÇÃO")
    add_para(document, [("O valor anual estimado é de R$ 1.180.116,40 (um milhão, cento e oitenta mil, cento e dezesseis reais e quarenta centavos), conforme a Nota Técnica de Justificativa de Preços, formada a partir do histórico das faturas das duas matrículas consumidoras.", RED, False, False, False)], style="Nível 02")
    add_para(document, [("A justificativa do preço observará as tarifas reguladas e as faturas de suporte. Não se aplica pesquisa competitiva de mercado, pois não há pluralidade de fornecedores para a área atendida; a estimativa e a justificativa do preço permanecem obrigatórias.", RED, False, False, False)], style="Nível 02")

    heading(document, "ADEQUAÇÃO ORÇAMENTÁRIA")
    optional(document, "As despesas correrão à conta da dotação orçamentária: [programa de trabalho, natureza da despesa, fonte, plano interno e declaração de disponibilidade].")

    heading(document, "ANEXO I — UNIDADES CONSUMIDORAS")
    optional(document, "Unidade 1: [endereço completo]; matrícula: [número]; economias: [quantidade].")
    optional(document, "Unidade 2: [endereço completo]; matrícula: [número]; economias: [quantidade].")

    add_para(document, [("Elaborado por: ", None, True, False, False), ("[nome, cargo, SIAPE e unidade]", BLUE, False, False, False)], after=230)
    add_para(document, [("Integrante requisitante/técnico: ", None, True, False, False), ("[nome, cargo, SIAPE e unidade]", BLUE, False, False, False)], after=230)
    add_para(document, [("Aprovado por: ", None, True, False, False), ("[autoridade competente]", BLUE, False, False, False)], after=50)
    document.save(output)
    return output


def build_enquadramento():
    output = OUT / "DOC ÁGUA - 06 — Formulário de Enquadramento da Inexigibilidade (base UFRRJ).docx"
    shutil.copy2(TEMPLATE, output)
    document = Document(output)
    wipe_body(document)
    add_cover(document, "FORMULÁRIO DE ENQUADRAMENTO — INEXIGIBILIDADE DE LICITAÇÃO")
    rationale(document, "Documento interno inspirado no Formulário de Enquadramento mencionado como documento 28 no ato autorizativo da UFRRJ. O formulário original não consta entre os arquivos publicados no PNCP.")

    heading(document, "IDENTIFICAÇÃO DA CONTRATAÇÃO")
    optional(document, "Processo administrativo: [preencher NUP].")
    add_para(document, [("Objeto: contratação da concessionária responsável pela área para o fornecimento contínuo de água tratada e a coleta/tratamento de esgoto sanitário nas unidades consumidoras do Complexo de São Cristóvão do Colégio Pedro II.", RED, False, False, False)], style="Nível 02")
    optional(document, "Contratada: [razão social e CNPJ a confirmar].")

    heading(document, "ENQUADRAMENTO LEGAL")
    add_para(document, [("A contratação direta será instruída como inexigibilidade de licitação, com fundamento no art. 74, inciso I, da Lei nº 14.133/2021, em razão da inviabilidade de competição decorrente da atuação exclusiva da concessionária na área atendida.", RED, False, False, False)], style="Nível 02")
    add_para(document, [("A instrução observará os elementos do art. 72 da Lei nº 14.133/2021, inclusive DFD, ETP quando cabível, análise de riscos, estimativa de despesa, parecer jurídico ou hipótese de parecer referencial aplicável, disponibilidade orçamentária, habilitação, razão da escolha, justificativa do preço e autorização da autoridade competente.", RED, False, False, False)], style="Nível 02")

    heading(document, "COMPROVAÇÃO DA INVIABILIDADE DE COMPETIÇÃO")
    add_para(document, [("Consta como documento territorial de suporte o Relatório de Caracterização da AGENERSA referente ao Bloco 4, que identifica São Cristóvão entre os bairros abrangidos pela área concedida à Águas do Rio 4.", RED, False, False, False)], style="Nível 02")
    optional(document, "Antes da conclusão e assinatura, anexar declaração atual da concessionária e/ou trecho do contrato de concessão que vincule expressamente cada matrícula e endereço do CPII à área de atendimento exclusiva.")
    optional(document, "Matrícula 1: [número] — endereço: [endereço completo]. Matrícula 2: [número] — endereço: [endereço completo].")
    rationale(document, "O relatório territorial da AGENERSA apoia a comprovação, mas não substitui o vínculo documental das unidades consumidoras específicas à concessionária.")

    heading(document, "RAZÃO DA ESCOLHA DA CONTRATADA")
    add_para(document, [("A escolha decorre da condição de concessionária responsável pela rede e pela prestação exclusiva dos serviços públicos de abastecimento de água e esgotamento sanitário na área das unidades consumidoras, não havendo alternativa competitiva para o objeto delimitado.", RED, False, False, False)], style="Nível 02")

    heading(document, "JUSTIFICATIVA DO PREÇO")
    add_para(document, [("A estimativa anual de R$ 1.180.116,40 será demonstrada pela Nota Técnica de Justificativa de Preços, memória de cálculo e faturas das duas matrículas. As tarifas são reguladas; por isso não há pesquisa competitiva de mercado, sem prejuízo da demonstração da razoabilidade do valor e do consumo histórico.", RED, False, False, False)], style="Nível 02")

    heading(document, "CONCLUSÃO")
    add_para(document, [("Após a juntada e a conferência dos documentos indicados neste formulário, conclui-se pelo enquadramento da contratação por inexigibilidade, nos termos do art. 74, inciso I, da Lei nº 14.133/2021, para o objeto e as unidades estritamente delimitados nos autos.", RED, False, False, False)], style="Nível 02")
    optional(document, "Local e data: [preencher].")
    add_para(document, [("Responsável pela instrução: ", None, True, False, False), ("[nome, cargo, SIAPE e unidade]", BLUE, False, False, False)], after=220)
    add_para(document, [("Área requisitante/técnica: ", None, True, False, False), ("[nome, cargo, SIAPE e unidade]", BLUE, False, False, False)], after=60)
    document.save(output)
    return output


if __name__ == "__main__":
    tr = build_tr()
    form = build_enquadramento()
    print(tr)
    print(form)
